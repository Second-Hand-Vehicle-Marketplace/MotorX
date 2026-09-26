"""Build the MotorX production guide from its editable Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'MotorX_Production_Security_and_Resilience_Guide.md'
OUTPUT = SOURCE.with_suffix('.docx')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.64)
section.left_margin = section.right_margin = Inches(0.72)
section.header_distance = section.footer_distance = Inches(0.27)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string('243447')
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.06
for name, size, color in [('Title', 34, '12334B'), ('Heading 1', 20, '12334B'), ('Heading 2', 13, '087D88')]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True
for name in ['List Bullet', 'List Number']:
    doc.styles[name].font.size = Pt(10.5)
    doc.styles[name].paragraph_format.space_after = Pt(5)
    doc.styles[name].paragraph_format.line_spacing = 1.04

header = section.header.paragraphs[0]
header.text = 'MOTORX  /  PRODUCTION ENGINEERING'
header.style = doc.styles['Caption']
header.runs[0].font.color.rgb = RGBColor.from_string('087D88')
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('EC2 security and resilience  •  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('64748B')

def inline(paragraph, text):
    for part in re.split(r'(\*\*.*?\*\*|https://\S+)', text):
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('https://'):
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(part, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '087D88'); props.append(color)
            size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
            run.append(props)
            text_node = OxmlElement('w:t'); text_node.text = part; run.append(text_node)
            link.append(run); paragraph._p.append(link)
        else:
            paragraph.add_run(part)

def add_table(rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = 'Light Shading Accent 1'
    table.autofit = False
    count = len(rows[0])
    widths = [1.30, 2.77, 2.76] if count == 3 else [1.40, 5.43]
    if count == 3 and rows[0][0] == 'ID': widths = [0.42, 2.81, 3.60]
    for col, width in zip(table.columns, widths): col.width = Inches(width)
    for i, row in enumerate(rows):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            inline(p, value)
            for run in p.runs: run.font.size = Pt(9)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '12334B'); cell._tc.get_or_add_tcPr().append(shade)
                for run in p.runs:
                    run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
        trpr = table.rows[i]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
first_page = True
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line == '@@page':
        doc.add_page_break(); first_page = False
    elif line.startswith('```'):
        chunk = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            chunk.append(lines[i]); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EFF5F7'); p._p.get_or_add_pPr().append(shade)
        run = p.add_run('\n'.join(chunk)); run.font.name = 'Consolas'; run.font.size = Pt(8.5)
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', x) for x in cells): rows.append(cells)
            i += 1
        add_table(rows); continue
    elif line.startswith('# '):
        p = doc.add_paragraph(line[2:], 'Title' if first_page else 'Heading 1')
        if first_page: p.paragraph_format.space_before = Pt(42)
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 2')
    elif line.startswith('- '):
        inline(doc.add_paragraph(style='List Bullet'), line[2:])
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        inline(p, line)
    else:
        p = doc.add_paragraph(); inline(p, line)
        if re.match(r'^\[\d+\]', line):
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs: run.font.size = Pt(9)
    i += 1

doc.core_properties.title = 'MotorX Production Security and Fault Tolerance'
doc.core_properties.subject = 'Practical security, reliable jobs and recovery for an EC2 deployment'
doc.core_properties.author = 'MotorX project documentation'
doc.core_properties.keywords = 'MotorX, EC2, security, fault tolerance, recovery, production'
doc.save(OUTPUT)
print(f'Created: {OUTPUT}')
print(f'Words in source: {len(SOURCE.read_text(encoding="utf-8").split())}')
print(f'Tables: {len(doc.tables)}; paragraphs: {len(doc.paragraphs)}')
